import os
import json
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog
from datetime import datetime, timedelta
import re
import shutil
import csv

# Для экспорта в PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Для экспорта в Excel
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# Для графика
import matplotlib.pyplot as plt
import numpy as np

# Путь к шрифту и иконке
ASSETS_DIR = Path(__file__).parent / "assets"
CHAKRA_FONT_PATH = ASSETS_DIR / "ChakraPetch-Regular.ttf"
ICON_PATH = Path(__file__).parent / "ico.ico"
USE_CHAKRA_FONT = CHAKRA_FONT_PATH.exists()

if USE_CHAKRA_FONT:
    try:
        pdfmetrics.registerFont(TTFont('ChakraPetch', str(CHAKRA_FONT_PATH)))
    except Exception as e:
        messagebox.showwarning("Шрифт", f"Не удалось загрузить шрифт ChakraPetch:\n{e}")
        USE_CHAKRA_FONT = False

# Файлы настроек и данных
SETTINGS_PATH = Path(__file__).parent / "settings.json"
AUDIT_LOG_PATH = None  # будет задан после загрузки base_dir


def load_base_dir():
    if SETTINGS_PATH.exists():
        try:
            with open(SETTINGS_PATH, 'r', encoding='utf-8') as f:
                settings = json.load(f)
                base_dir = Path(settings.get("base_dir", ""))
                if base_dir.is_dir():
                    return base_dir
        except Exception:
            pass

    root = tk.Tk()
    root.withdraw()
    messagebox.showinfo("Первый запуск", "Пожалуйста, выберите папку с файлами base.json и solutor.json.")
    folder = filedialog.askdirectory(title="Выберите папку с базой данных")
    root.destroy()

    if not folder:
        messagebox.showerror("Ошибка", "Папка не выбрана. Программа будет закрыта.")
        exit()

    base_dir = Path(folder)
    try:
        save_base_dir(base_dir)
    except Exception:
        pass
    return base_dir


def save_base_dir(base_dir: Path):
    with open(SETTINGS_PATH, 'w', encoding='utf-8') as f:
        json.dump({"base_dir": str(base_dir)}, f, ensure_ascii=False, indent=4)


def ensure_base_exists(base_path: Path):
    base_path.parent.mkdir(parents=True, exist_ok=True)
    if not base_path.exists():
        with open(base_path, 'w', encoding='utf-8') as f:
            json.dump([], f, ensure_ascii=False, indent=4)


def ensure_solutor_exists(solutor_path: Path):
    solutor_path.parent.mkdir(parents=True, exist_ok=True)
    if not solutor_path.exists():
        default_payers = ["ИТ", "Бухгалтерия", "Отдел закупок", "Дирекция"]
        with open(solutor_path, 'w', encoding='utf-8') as f:
            json.dump(default_payers, f, ensure_ascii=False, indent=4)


def log_action(action: str):
    """Записывает действие в audit.log"""
    if AUDIT_LOG_PATH is None:
        return
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    try:
        with open(AUDIT_LOG_PATH, 'a', encoding='utf-8') as f:
            f.write(f"[{timestamp}] {action}\n")
    except Exception:
        pass  # Не критично, если лог не пишется


def validate_date(date_str: str) -> bool:
    if not re.fullmatch(r'\d{2}\.\d{2}\.\d{4}', date_str):
        return False
    try:
        datetime.strptime(date_str, "%d.%m.%Y")
        return True
    except ValueError:
        return False


def validate_amount(amount_str: str) -> bool:
    if not amount_str.strip():
        return False
    try:
        float(amount_str.replace(" ", "").replace(",", "."))
        return True
    except ValueError:
        return False


class PayersManager:
    def __init__(self, parent, solutor_path: Path, readonly_mode: bool):
        self.parent = parent
        self.solutor_path = solutor_path
        self.readonly_mode = readonly_mode
        self.load_payers()

    def load_payers(self):
        if not self.solutor_path.exists():
            self.payers = []
            return
        try:
            with open(self.solutor_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                if isinstance(data, list):
                    self.payers = [str(p).strip() for p in data if str(p).strip()]
                else:
                    self.payers = []
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить список плательщиков:\n{e}")
            self.payers = []

    def save_payers(self):
        if self.readonly_mode:
            return
        try:
            with open(self.solutor_path, 'w', encoding='utf-8') as f:
                json.dump(self.payers, f, ensure_ascii=False, indent=4)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить список плательщиков:\n{e}")

    def open_payers_window(self):
        win = tk.Toplevel(self.parent)
        win.title("Управление плательщиками")
        win.geometry("500x400")
        win.grab_set()

        listbox = tk.Listbox(win, font=("Arial", 10), selectmode=tk.SINGLE)
        listbox.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        for payer in self.payers:
            listbox.insert(tk.END, payer)

        def add_payer():
            new_payer = simpledialog.askstring("Новый плательщик", "Введите название плательщика:")
            if new_payer and new_payer.strip():
                new_payer = new_payer.strip()
                if new_payer not in self.payers:
                    self.payers.append(new_payer)
                    self.save_payers()
                    log_action(f"Добавлен плательщик: {new_payer}")
                    listbox.insert(tk.END, new_payer)
                else:
                    messagebox.showinfo("Инфо", "Такой плательщик уже существует.")

        def delete_payer():
            sel = listbox.curselection()
            if not sel:
                return
            idx = sel[0]
            payer = self.payers[idx]
            if messagebox.askyesno("Подтверждение", f"Удалить плательщика '{payer}'?"):
                del self.payers[idx]
                self.save_payers()
                log_action(f"Удалён плательщик: {payer}")
                listbox.delete(idx)

        btn_frame = tk.Frame(win)
        btn_frame.pack(pady=5)

        tk.Button(btn_frame, text="Добавить", command=add_payer, width=12).pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="Удалить", command=delete_payer, width=12).pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="Закрыть", command=win.destroy, width=12).pack(side=tk.LEFT, padx=5)

        if self.readonly_mode:
            for widget in btn_frame.winfo_children():
                widget.config(state='disabled')


class RegistrumApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Registrum — Реестр счетов покупок")
        self.root.state('zoomed')

        # Установка иконки, если существует
        if ICON_PATH.exists():
            try:
                self.root.iconbitmap(str(ICON_PATH))
            except Exception:
                pass

        self.base_dir = load_base_dir()
        global AUDIT_LOG_PATH
        AUDIT_LOG_PATH = self.base_dir / "audit.log"

        self.base_path = self.base_dir / "base.json"
        self.solutor_path = self.base_dir / "solutor.json"

        self.readonly_mode = not self.can_write_to_base_dir()
        if self.readonly_mode:
            self.root.title("Registrum — Реестр счетов покупок [Только чтение]")

        try:
            ensure_base_exists(self.base_path)
            ensure_solutor_exists(self.solutor_path)
        except Exception as e:
            if not self.readonly_mode:
                messagebox.showerror("Ошибка", f"Не удалось создать файлы:\n{e}")
                exit()

        # Автоматическое резервное копирование (раз в 24 часа)
        self.auto_backup()

        self.payers_manager = PayersManager(root, self.solutor_path, self.readonly_mode)
        self.payers_manager.load_payers()

        # === Верхняя панель ===
        top_frame = tk.Frame(root)
        top_frame.pack(pady=5, padx=20, fill=tk.X)

        # Поиск
        tk.Label(top_frame, text="Поиск:", font=("Arial", 10)).pack(side=tk.LEFT)
        self.search_var = tk.StringVar()
        self.search_var.trace("w", self.on_search_change)
        self.search_entry = tk.Entry(top_frame, textvariable=self.search_var, font=("Arial", 10), width=30)
        self.search_entry.pack(side=tk.LEFT, padx=(5, 10))

        # Фильтр по дате
        tk.Label(top_frame, text="Период:", font=("Arial", 10)).pack(side=tk.LEFT, padx=(10, 5))
        self.date_from_var = tk.StringVar()
        self.date_to_var = tk.StringVar()
        self.date_from_entry = tk.Entry(top_frame, textvariable=self.date_from_var, font=("Arial", 10), width=10)
        self.date_to_entry = tk.Entry(top_frame, textvariable=self.date_to_var, font=("Arial", 10), width=10)
        self.date_from_entry.pack(side=tk.LEFT)
        tk.Label(top_frame, text="–", font=("Arial", 10)).pack(side=tk.LEFT)
        self.date_to_entry.pack(side=tk.LEFT, padx=(0, 10))
        tk.Button(top_frame, text="Применить", command=self.apply_date_filter, font=("Arial", 9)).pack(side=tk.LEFT)

        tk.Button(top_frame, text="Очистить", command=self.clear_filters).pack(side=tk.LEFT, padx=(10, 5))
        self.btn_backup = tk.Button(top_frame, text="Резерв", command=self.create_backup)
        self.btn_backup.pack(side=tk.LEFT, padx=(0, 20))

        self.status_label = tk.Label(top_frame, text="", font=("Arial", 10, "bold"))
        self.status_label.pack(side=tk.LEFT)

        # Кнопки справа
        btn_frame = tk.Frame(top_frame)
        btn_frame.pack(side=tk.RIGHT)

        self.btn_pdf = tk.Button(btn_frame, text="В PDF", command=self.export_to_pdf, width=12, height=1)
        self.btn_excel = tk.Button(btn_frame, text="В Excel", command=self.export_to_excel, width=12, height=1)
        self.btn_import = tk.Button(btn_frame, text="Импорт CSV", command=self.import_from_csv, width=12, height=1)
        self.btn_payers = tk.Button(btn_frame, text="Плательщики", command=self.open_payers_window, width=12, height=1)
        self.btn_settings = tk.Button(btn_frame, text="Настройки", command=self.open_settings, width=12, height=1)
        self.btn_chart = tk.Button(btn_frame, text="График", command=self.show_chart, width=12, height=1)
        self.btn_info = tk.Button(btn_frame, text="Инфо", command=self.show_info, width=12, height=1)
        self.btn_exit = tk.Button(btn_frame, text="Выход", command=self.on_exit, width=12, height=1)

        buttons = [self.btn_pdf, self.btn_excel, self.btn_import, self.btn_payers, self.btn_settings, self.btn_chart, self.btn_info, self.btn_exit]
        for i, btn in enumerate(buttons):
            btn.grid(row=0, column=i, padx=2)

        if self.readonly_mode:
            self.btn_backup.config(state='disabled')
            self.btn_settings.config(state='disabled')
            self.btn_payers.config(state='disabled')
            self.btn_import.config(state='disabled')

        # Таблица
        table_frame = tk.Frame(root)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        self.columns = ["Дата", "Заказ", "Сумма", "Поставщик", "Плательщик", "Инициатор", "Обоснование", "Оплата", "Забрал", "Комментарии"]
        self.tree = ttk.Treeview(table_frame, columns=self.columns, show='headings')

        for col in self.columns:
            self.tree.heading(col, text=col, command=lambda c=col: self.sort_column(c))
            self.tree.column(col, width=100, anchor='center')

        v_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        h_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscroll=v_scroll.set, xscroll=h_scroll.set)

        self.tree.grid(row=0, column=0, sticky='nsew')
        v_scroll.grid(row=0, column=1, sticky='ns')
        h_scroll.grid(row=1, column=0, sticky='ew')

        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        # Контекстное меню
        self.context_menu = tk.Menu(self.tree, tearoff=0)
        self.context_menu.add_command(label="Удалить", command=self.delete_selected)
        self.tree.bind("<Button-3>", self.show_context_menu)
        self.tree.bind("<Double-1>", self.on_double_click)

        if self.readonly_mode:
            self.context_menu.delete(0)

        # Форма ввода
        form_frame = tk.Frame(root)
        form_frame.pack(pady=10, padx=20, fill=tk.X)

        tk.Label(form_frame, text="Форма добавления / редактирования заказа:", font=("Arial", 12, "bold")).pack(anchor='w')

        self.entries = {}
        grid_frame = tk.Frame(form_frame)
        grid_frame.pack(pady=5, fill=tk.X)

        short_fields = ["Дата", "Заказ", "Сумма", "Поставщик", "Инициатор", "Оплата", "Забрал"]
        long_fields = ["Обоснование", "Комментарии"]

        for i, field in enumerate(short_fields):
            row = i // 2
            col = (i % 2) * 2
            tk.Label(grid_frame, text=f"{field}:", font=("Arial", 10)).grid(row=row, column=col, sticky='e', padx=5, pady=3)
            entry = tk.Entry(grid_frame, font=("Arial", 10), width=30)
            entry.grid(row=row, column=col+1, sticky='w', padx=5, pady=3)
            self.entries[field] = entry

        # Поле "Плательщик" — Combobox
        row = len(short_fields) // 2
        col = 0
        tk.Label(grid_frame, text="Плательщик:", font=("Arial", 10)).grid(row=row, column=col, sticky='e', padx=5, pady=3)
        self.payer_combobox = ttk.Combobox(grid_frame, values=self.payers_manager.payers, font=("Arial", 10), width=28, state="readonly")
        self.payer_combobox.grid(row=row, column=col+1, sticky='w', padx=5, pady=3)
        self.entries["Плательщик"] = self.payer_combobox

        for field in long_fields:
            tk.Label(grid_frame, text=f"{field}:", font=("Arial", 10)).grid(
                row=row + 1 + long_fields.index(field), column=0, sticky='ne', padx=5, pady=5)
            text = tk.Text(grid_frame, font=("Arial", 10), width=60, height=3)
            text.grid(row=row + 1 + long_fields.index(field), column=1, columnspan=3, sticky='ew', padx=5, pady=5)
            self.entries[field] = text

        grid_frame.grid_columnconfigure(1, weight=1)
        grid_frame.grid_columnconfigure(3, weight=1)

        form_btn_frame = tk.Frame(form_frame)
        form_btn_frame.pack(pady=10)

        self.btn_new = tk.Button(form_btn_frame, text="Новый", command=self.clear_form, width=20, height=2)
        self.btn_save_order = tk.Button(form_btn_frame, text="Сохранить заказ", command=self.save_order, width=20, height=2)
        self.btn_cancel = tk.Button(form_btn_frame, text="Отмена", command=self.clear_form, width=20, height=2)

        self.btn_new.pack(side=tk.LEFT, padx=10)
        self.btn_save_order.pack(side=tk.LEFT, padx=10)
        self.btn_cancel.pack(side=tk.LEFT, padx=10)

        if self.readonly_mode:
            self.btn_new.config(state='disabled')
            self.btn_save_order.config(state='disabled')
            self.btn_cancel.config(state='disabled')
            for widget in self.entries.values():
                if isinstance(widget, tk.Entry):
                    widget.config(state='disabled')
                elif isinstance(widget, tk.Text):
                    widget.config(state='disabled')
                elif isinstance(widget, ttk.Combobox):
                    widget.config(state='disabled')

        self.editing_index = None
        self.all_data = []
        self.filtered_data = []  # данные после применения фильтров
        self.load_table()
        self.clear_form()
        self.root.protocol("WM_DELETE_WINDOW", self.on_exit)
        self.update_yearly_total()

    def can_write_to_base_dir(self) -> bool:
        if not self.base_dir:
            return False
        try:
            test_file = self.base_dir / ".write_test_registrum"
            test_file.write_text("ok", encoding='utf-8')
            test_file.unlink()
            return True
        except (OSError, IOError, PermissionError):
            return False

    def auto_backup(self):
        """Создаёт резервную копию, если прошло более 24 часов с последней."""
        if self.readonly_mode:
            return

        backup_marker = self.base_dir / "last_backup.txt"
        now = datetime.now()
        should_backup = True

        if backup_marker.exists():
            try:
                with open(backup_marker, 'r') as f:
                    last_time_str = f.read().strip()
                    last_time = datetime.fromisoformat(last_time_str)
                    if (now - last_time) < timedelta(hours=24):
                        should_backup = False
            except Exception:
                pass

        if should_backup:
            self.create_backup(silent=True)
            try:
                with open(backup_marker, 'w') as f:
                    f.write(now.isoformat())
            except Exception:
                pass

    def create_backup(self, silent=False):
        if self.readonly_mode:
            if not silent:
                messagebox.showwarning("Доступ запрещён", "Режим только для чтения.")
            return
        if not self.base_path.exists():
            if not silent:
                messagebox.showwarning("Предупреждение", "Файл базы не существует!")
            return

        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        backup_base_path = self.base_dir / f"base_{timestamp}.json"
        backup_solutor_path = self.base_dir / f"solutor_{timestamp}.json"
        try:
            shutil.copy2(self.base_path, backup_base_path)
            if self.solutor_path.exists():
                shutil.copy2(self.solutor_path, backup_solutor_path)
            log_action(f"Созданы резервные копии: {backup_base_path.name}, {backup_solutor_path.name}")
            if not silent:
                messagebox.showinfo("Успех", f"Резервные копии созданы:\n{backup_base_path}\n{backup_solutor_path}")
        except Exception as e:
            if not silent:
                messagebox.showerror("Ошибка", f"Не удалось создать резервные копии:\n{e}")

    def load_data(self):
        if not self.base_path.exists():
            return []
        try:
            with open(self.base_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить базу:\n{e}")
            return []

    def save_data(self, data):
        if self.readonly_mode:
            return
        try:
            with open(self.base_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, ensure_ascii=False, indent=4)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить базу:\n{e}")

    def load_table(self):
        self.all_data = self.load_data()
        self.sort_by_date_desc()
        self.apply_filters()
        self.auto_adjust_column_widths()

    def sort_by_date_desc(self):
        def _sort_key(record):
            date_str = record.get("Дата", "")
            if re.fullmatch(r'\d{2}\.\d{2}\.\d{4}', str(date_str)):
                try:
                    d = datetime.strptime(date_str, "%d.%m.%Y")
                    return (d.year, d.month, d.day)
                except:
                    pass
            return (0, 0, 0)
        self.all_data.sort(key=_sort_key, reverse=True)

    def apply_filters(self):
        """Применяет поиск и фильтр по дате."""
        search_term = self.search_var.get().lower()
        date_from = self.date_from_var.get().strip()
        date_to = self.date_to_var.get().strip()

        filtered = []
        for record in self.all_data:
            # Поиск
            if search_term:
                match = any(search_term in str(record.get(col, "")).lower() for col in self.columns)
                if not match:
                    continue

            # Фильтр по дате
            date_str = record.get("Дата", "").strip()
            if date_from or date_to:
                if not validate_date(date_str):
                    continue
                record_date = datetime.strptime(date_str, "%d.%m.%Y")
                if date_from:
                    if not validate_date(date_from):
                        continue
                    from_date = datetime.strptime(date_from, "%d.%m.%Y")
                    if record_date < from_date:
                        continue
                if date_to:
                    if not validate_date(date_to):
                        continue
                    to_date = datetime.strptime(date_to, "%d.%m.%Y")
                    if record_date > to_date:
                        continue

            filtered.append(record)

        self.filtered_data = filtered
        self.refresh_table_view()

    def apply_date_filter(self):
        self.apply_filters()

    def clear_filters(self):
        self.search_var.set("")
        self.date_from_var.set("")
        self.date_to_var.set("")
        self.apply_filters()

    def on_search_change(self, *args):
        self.apply_filters()

    def refresh_table_view(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        for record in self.filtered_data:
            values = [record.get(col, "") for col in self.columns]
            self.tree.insert('', tk.END, values=values)

        self.update_yearly_total()

    def auto_adjust_column_widths(self):
        default_widths = {
            "Дата": 80,
            "Заказ": 100,
            "Сумма": 90,
            "Поставщик": 130,
            "Плательщик": 120,
            "Инициатор": 90,
            "Оплата": 60,
            "Забрал": 60,
            "Обоснование": 400,
            "Комментарии": 250,
        }
        for col in self.columns:
            self.tree.column(col, width=default_widths.get(col, 120))

    def clear_form(self):
        if self.readonly_mode:
            return
        self.editing_index = None
        for field in ["Дата", "Заказ", "Сумма", "Поставщик", "Инициатор", "Оплата", "Забрал"]:
            self.entries[field].delete(0, tk.END)
            if field == "Дата":
                self.entries[field].insert(0, self.get_today())
            elif field == "Инициатор":
                self.entries[field].insert(0, "ИТ")
            elif field in ("Оплата", "Забрал"):
                self.entries[field].insert(0, "Да")
        self.payer_combobox.set("")
        for field in ["Обоснование", "Комментарии"]:
            self.entries[field].delete("1.0", tk.END)

    def get_today(self):
        return datetime.now().strftime("%d.%m.%Y")

    def on_double_click(self, event):
        selected = self.tree.selection()
        if not selected:
            return
        item = selected[0]
        # Найти индекс в filtered_data
        index_in_filtered = self.tree.index(item)
        if index_in_filtered >= len(self.filtered_data):
            return
        record = self.filtered_data[index_in_filtered]

        # Найти оригинальный индекс в all_data
        try:
            self.editing_index = self.all_data.index(record)
        except ValueError:
            self.editing_index = None
            return

        for field in ["Дата", "Заказ", "Сумма", "Поставщик", "Инициатор", "Оплата", "Забрал"]:
            self.entries[field].delete(0, tk.END)
            self.entries[field].insert(0, record.get(field, ""))
        payer = record.get("Плательщик", "")
        if payer in self.payers_manager.payers:
            self.payer_combobox.set(payer)
        else:
            self.payer_combobox.set("")
        for field in ["Обоснование", "Комментарии"]:
            self.entries[field].delete("1.0", tk.END)
            self.entries[field].insert("1.0", record.get(field, ""))

    def save_order(self):
        if self.readonly_mode:
            messagebox.showwarning("Доступ запрещён", "Режим только для чтения.")
            return

        # Валидация
        date_val = self.entries["Дата"].get().strip()
        if not validate_date(date_val):
            messagebox.showerror("Ошибка", "Некорректная дата. Используйте формат ДД.ММ.ГГГГ.")
            return

        sum_val = self.entries["Сумма"].get().strip()
        if not validate_amount(sum_val):
            messagebox.showerror("Ошибка", "Сумма должна быть числом.")
            return

        supplier = self.entries["Поставщик"].get().strip()
        if not supplier:
            messagebox.showerror("Ошибка", "Поле 'Поставщик' обязательно для заполнения.")
            return

        record = {}
        for field in ["Дата", "Заказ", "Сумма", "Поставщик", "Инициатор", "Оплата", "Забрал"]:
            record[field] = self.entries[field].get().strip()
        record["Плательщик"] = self.payer_combobox.get().strip()
        for field in ["Обоснование", "Комментарии"]:
            record[field] = self.entries[field].get("1.0", tk.END).strip()

        data = self.all_data
        if self.editing_index is not None and 0 <= self.editing_index < len(data):
            old_record = data[self.editing_index]
            data[self.editing_index] = record
            self.save_data(data)
            log_action(f"Обновлена запись: {record.get('Заказ', 'без номера')}")
            self.load_table()
            messagebox.showinfo("Успех", "Запись успешно обновлена!")
        else:
            data.append(record)
            self.save_data(data)
            log_action(f"Добавлена запись: {record.get('Заказ', 'без номера')} на {record.get('Сумма', '0')} руб.")
            self.load_table()
            messagebox.showinfo("Успех", "Новый заказ успешно добавлен!")

        self.clear_form()

    def export_to_pdf(self):
        data = self.filtered_data
        if not data:
            messagebox.showwarning("Предупреждение", "Нет данных для экспорта!")
            return

        file_path = filedialog.asksaveasfilename(
            title="Сохранить как PDF",
            defaultextension=".pdf",
            filetypes=[("PDF files", "*.pdf")]
        )
        if not file_path:
            return

        try:
            doc = SimpleDocTemplate(file_path, pagesize=landscape(A4), leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
            elements = []

            styles = getSampleStyleSheet()
            title_style = styles['Title'].clone('CustomTitle')
            if USE_CHAKRA_FONT:
                title_style.fontName = 'ChakraPetch'
            title_style.fontSize = 16
            title = Paragraph("Реестр счетов покупок (Registrum)", title_style)
            elements.append(title)
            elements.append(Spacer(1, 12))

            # Подсчёт итога
            total_sum = 0.0
            for record in data:
                try:
                    s = record.get("Сумма", "0").replace(" ", "").replace(",", ".")
                    total_sum += float(s)
                except:
                    pass
            total_paragraph = Paragraph(f"<b>Итого: {int(total_sum):,} руб.</b>".replace(',', ' '), styles['Normal'])
            elements.append(total_paragraph)
            elements.append(Spacer(1, 12))

            # Таблица
            table_data = [self.columns]
            for record in data:
                row = [str(record.get(col, "")) for col in self.columns]
                table_data.append(row)

            total_width = landscape(A4)[0] - 72
            col_widths = []
            for col in self.columns:
                if col == "Обоснование":
                    col_widths.append(total_width * 0.30)
                elif col in ("Комментарии", "Поставщик", "Плательщик"):
                    col_widths.append(total_width * 0.12)
                else:
                    col_widths.append(total_width * 0.07)

            actual_sum = sum(col_widths)
            if actual_sum > 0:
                col_widths = [w * total_width / actual_sum for w in col_widths]

            table_style = [
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'ChakraPetch' if USE_CHAKRA_FONT else 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('FONTSIZE', (0, 1), (-1, -1), 7),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ]
            if USE_CHAKRA_FONT:
                table_style.append(('FONTNAME', (0, 1), (-1, -1), 'ChakraPetch'))

            table = Table(table_data, colWidths=col_widths, repeatRows=1)
            table.setStyle(TableStyle(table_style))
            elements.append(table)

            # Нумерация страниц
            def add_page_number(canvas, doc):
                page_num = canvas.getPageNumber()
                text = f"Стр. {page_num}"
                canvas.setFont('Helvetica', 9)
                canvas.drawRightString(landscape(A4)[0] - 36, 20, text)

            doc.build(elements, onFirstPage=add_page_number, onLaterPages=add_page_number)
            messagebox.showinfo("Успех", f"Файл PDF сохранён:\n{file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать PDF:\n{e}")

    def export_to_excel(self):
        data = self.filtered_data
        if not data:
            messagebox.showwarning("Предупреждение", "Нет данных для экспорта!")
            return

        file_path = filedialog.asksaveasfilename(
            title="Сохранить как Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if not file_path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Реестр счетов"
            ws.append(self.columns)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for record in data:
                row = [record.get(col, "") for col in self.columns]
                ws.append(row)

            # Итоговая строка
            total_sum = 0.0
            for record in data:
                try:
                    s = record.get("Сумма", "0").replace(" ", "").replace(",", ".")
                    total_sum += float(s)
                except:
                    pass
            ws.append([""] * (len(self.columns) - 1) + [f"Итого: {int(total_sum):,} руб.".replace(',', ' ')])

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 40)
                ws.column_dimensions[column].width = adjusted_width

            wb.save(file_path)
            messagebox.showinfo("Успех", f"Файл Excel сохранён:\n{file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать Excel:\n{e}")

    def import_from_csv(self):
        if self.readonly_mode:
            messagebox.showwarning("Доступ запрещён", "Режим только для чтения.")
            return

        file_path = filedialog.askopenfilename(
            title="Выберите CSV-файл для импорта",
            filetypes=[("CSV files", "*.csv")]
        )
        if not file_path:
            return

        try:
            new_records = []
            with open(file_path, 'r', encoding='utf-8-sig') as f:
                reader = csv.DictReader(f, delimiter=';')
                for row in reader:
                    # Проверяем обязательные поля
                    if not row.get("Поставщик") or not row.get("Сумма"):
                        continue
                    # Приводим к нужному формату
                    record = {col: row.get(col, "") for col in self.columns}
                    new_records.append(record)

            if not new_records:
                messagebox.showwarning("Предупреждение", "Нет корректных записей для импорта.")
                return

            if messagebox.askyesno("Подтверждение", f"Будет добавлено {len(new_records)} записей. Продолжить?"):
                self.all_data.extend(new_records)
                self.save_data(self.all_data)
                log_action(f"Импортировано {len(new_records)} записей из CSV")
                self.load_table()
                messagebox.showinfo("Успех", "Импорт завершён успешно!")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось импортировать CSV:\n{e}")

    def show_info(self):
        info_text = (
            "Registrum v0.7\n\n"
            "Система учёта счетов покупок\n"
            "Автор: Разин Г.В.\n"
            "© 2025"
        )
        messagebox.showinfo("О программе", info_text)

    def open_settings(self):
        if self.readonly_mode:
            messagebox.showwarning("Доступ запрещён", "Режим только для чтения.")
            return

        settings_win = tk.Toplevel(self.root)
        settings_win.title("Настройки")
        settings_win.geometry("600x150")
        settings_win.resizable(False, False)
        settings_win.grab_set()

        tk.Label(settings_win, text="Текущая папка базы данных:", font=("Arial", 10)).pack(pady=5)
        current_path = tk.Entry(settings_win, font=("Arial", 10), width=70)
        current_path.insert(0, str(self.base_dir))
        current_path.config(state='readonly')
        current_path.pack(pady=5)

        def change_path():
            new_dir = filedialog.askdirectory(title="Выберите папку для хранения базы данных")
            if not new_dir:
                return
            new_dir_path = Path(new_dir)
            try:
                new_dir_path.mkdir(parents=True, exist_ok=True)
                test_file = new_dir_path / ".test_write"
                test_file.write_text("ok", encoding='utf-8')
                test_file.unlink()
            except Exception as e:
                messagebox.showerror("Ошибка", f"Нет прав на запись:\n{e}")
                return

            self.base_dir = new_dir_path
            global AUDIT_LOG_PATH
            AUDIT_LOG_PATH = self.base_dir / "audit.log"
            self.base_path = new_dir_path / "base.json"
            self.solutor_path = new_dir_path / "solutor.json"
            save_base_dir(self.base_dir)
            ensure_base_exists(self.base_path)
            ensure_solutor_exists(self.solutor_path)
            self.payers_manager.solutor_path = self.solutor_path
            self.payers_manager.load_payers()
            self.payer_combobox['values'] = self.payers_manager.payers
            self.load_table()
            self.clear_form()
            current_path.config(state='normal')
            current_path.delete(0, tk.END)
            current_path.insert(0, str(self.base_dir))
            current_path.config(state='readonly')
            messagebox.showinfo("Успех", f"Папка базы изменена на:\n{self.base_dir}")

        tk.Button(settings_win, text="Изменить путь", command=change_path, width=20).pack(pady=10)

    def open_payers_window(self):
        self.payers_manager.open_payers_window()

    def show_context_menu(self, event):
        item = self.tree.identify_row(event.y)
        if item:
            self.tree.selection_set(item)
            self.context_menu.post(event.x_root, event.y_root)

    def delete_selected(self):
        if self.readonly_mode:
            messagebox.showwarning("Доступ запрещён", "Режим только для чтения.")
            return

        selected = self.tree.selection()
        if not selected:
            return
        if not messagebox.askyesno("Подтверждение", "Удалить выбранную запись?"):
            return

        index_in_filtered = self.tree.index(selected[0])
        if index_in_filtered >= len(self.filtered_data):
            return
        record_to_delete = self.filtered_data[index_in_filtered]

        try:
            self.all_data.remove(record_to_delete)
        except ValueError:
            messagebox.showerror("Ошибка", "Запись не найдена в базе.")
            return

        self.save_data(self.all_data)
        log_action(f"Удалена запись: {record_to_delete.get('Заказ', 'без номера')}")
        self.load_table()
        self.clear_form()
        messagebox.showinfo("Успех", "Запись удалена.")

    def sort_column(self, col):
        reverse = False
        if hasattr(self, '_last_sorted_col') and self._last_sorted_col == col:
            reverse = not getattr(self, '_last_sorted_reverse', False)
        else:
            reverse = False

        self._last_sorted_col = col
        self._last_sorted_reverse = reverse

        def _sort_key(record):
            val = record.get(col, "")
            if col == "Дата":
                if re.fullmatch(r'\d{2}\.\d{2}\.\d{4}', str(val)):
                    try:
                        d = datetime.strptime(val, "%d.%m.%Y")
                        return (d.year, d.month, d.day)
                    except:
                        pass
                return (9999, 99, 99)
            else:
                try:
                    return float(str(val).replace(" ", "").replace(",", "."))
                except:
                    return str(val).lower()

        self.all_data.sort(key=_sort_key, reverse=reverse)
        self.apply_filters()

    def update_yearly_total(self):
        current_year = datetime.now().year
        total_current = 0.0
        total_all = 0.0
        for record in self.all_data:
            date_str = record.get("Дата", "").strip()
            sum_str = record.get("Сумма", "").strip()

            year = None
            if re.match(r'\d{2}\.\d{2}\.\d{4}', date_str):
                try:
                    year = datetime.strptime(date_str, "%d.%m.%Y").year
                except:
                    pass

            try:
                sum_val = float(sum_str.replace(" ", "").replace(",", "."))
                total_all += sum_val
                if year == current_year:
                    total_current += sum_val
            except:
                pass

        self.status_label.config(
            text=f"Текущий год: {int(total_current):,} руб. | Всего: {int(total_all):,} руб.".replace(',', ' ')
        )

    def show_chart(self):
        data = self.filtered_data
        if not data:
            messagebox.showwarning("Предупреждение", "Нет данных для построения графика!")
            return

        chart_win = tk.Toplevel(self.root)
        chart_win.title("Выбор типа графика")
        chart_win.geometry("350x270")
        chart_win.grab_set()

        tk.Label(chart_win, text="Выберите тип графика:", font=("Arial", 12, "bold")).pack(pady=(15, 10))

        chart_type = tk.StringVar(value="yearly_total")

        tk.Radiobutton(chart_win, text="Общий по годам", variable=chart_type, value="yearly_total", font=("Arial", 10)).pack(anchor=tk.W, padx=30, pady=2)
        tk.Radiobutton(chart_win, text="Сравнение плательщиков по годам", variable=chart_type, value="payer_comparison", font=("Arial", 10)).pack(anchor=tk.W, padx=30, pady=2)
        tk.Radiobutton(chart_win, text="Детализация по годам", variable=chart_type, value="yearly_detail", font=("Arial", 10)).pack(anchor=tk.W, padx=30, pady=2)
        tk.Radiobutton(chart_win, text="По месяцам (текущий год)", variable=chart_type, value="monthly_current", font=("Arial", 10)).pack(anchor=tk.W, padx=30, pady=2)

        def show_selected_chart():
            chart_win.destroy()
            if chart_type.get() == "yearly_total":
                self._show_yearly_total_chart(data)
            elif chart_type.get() == "payer_comparison":
                self._show_payer_comparison_chart(data)
            elif chart_type.get() == "yearly_detail":
                self._show_yearly_detail_chart(data)
            elif chart_type.get() == "monthly_current":
                self._show_monthly_chart(data)

        tk.Button(chart_win, text="Показать график", command=show_selected_chart, font=("Arial", 10), width=20).pack(pady=20)

    def _show_yearly_total_chart(self, data):
        yearly_totals = {}
        for record in data:
            date_str = record.get("Дата", "").strip()
            sum_str = record.get("Сумма", "").strip()
            year = None
            if re.match(r'\d{2}\.\d{2}\.\d{4}', date_str):
                try:
                    year = datetime.strptime(date_str, "%d.%m.%Y").year
                except:
                    continue
            try:
                sum_val = float(sum_str.replace(" ", "").replace(",", "."))
            except:
                continue
            if year not in yearly_totals:
                yearly_totals[year] = 0.0
            yearly_totals[year] += sum_val

        if not yearly_totals:
            messagebox.showwarning("Предупреждение", "Не удалось извлечь данные для графика!")
            return

        years = sorted(yearly_totals.keys())
        totals = [yearly_totals[year] for year in years]

        fig, ax = plt.subplots(figsize=(12, 7))
        bars = ax.bar(years, totals, color='steelblue')
        for bar in bars:
            height = bar.get_height()
            ax.annotate(f'{int(height):,}'.replace(',', ' '),
                        xy=(bar.get_x() + bar.get_width() / 2, height),
                        xytext=(0, 3),
                        textcoords="offset points",
                        ha='center', va='bottom', fontweight='bold')

        ax.set_title("Общие затраты по годам", fontsize=16, fontweight='bold')
        ax.set_xlabel("Год", fontsize=12)
        ax.set_ylabel("Сумма (руб.)", fontsize=12)
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f'{int(x):,}'.replace(',', ' ')))
        ax.set_xticks(years)
        ax.set_xticklabels([str(y) for y in years])
        ax.grid(axis='y', linestyle='--', alpha=0.7)
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.set_facecolor('#f8f9fa')

        manager = plt.get_current_fig_manager()
        try:
            manager.window.state('zoomed')
        except:
            pass
        plt.tight_layout()
        plt.show()

    def _show_payer_comparison_chart(self, data):
        yearly_payer_totals = {}
        for record in data:
            date_str = record.get("Дата", "").strip()
            sum_str = record.get("Сумма", "").strip()
            payer = record.get("Плательщик", "").strip()
            if not payer:
                continue
            year = None
            if re.match(r'\d{2}\.\d{2}\.\d{4}', date_str):
                try:
                    year = datetime.strptime(date_str, "%d.%m.%Y").year
                except:
                    continue
            try:
                sum_val = float(sum_str.replace(" ", "").replace(",", "."))
            except:
                continue
            if year not in yearly_payer_totals:
                yearly_payer_totals[year] = {}
            if payer not in yearly_payer_totals[year]:
                yearly_payer_totals[year][payer] = 0.0
            yearly_payer_totals[year][payer] += sum_val

        if not yearly_payer_totals:
            messagebox.showwarning("Предупреждение", "Не удалось извлечь данные для графика!")
            return

        years = sorted(yearly_payer_totals.keys())
        all_payers = sorted(set(payer for year_data in yearly_payer_totals.values() for payer in year_data.keys()))
        payer_totals_per_year = {payer: [] for payer in all_payers}
        for year in years:
            for payer in all_payers:
                payer_totals_per_year[payer].append(yearly_payer_totals[year].get(payer, 0.0))

        fig, ax = plt.subplots(figsize=(14, 8))
        bottom = np.zeros(len(years))
        colors = plt.cm.tab20(np.linspace(0, 1, len(all_payers)))
        for i, payer in enumerate(all_payers):
            ax.bar(years, payer_totals_per_year[payer], bottom=bottom, label=payer, color=colors[i])
            bottom += np.array(payer_totals_per_year[payer])

        for i, year in enumerate(years):
            total = sum(payer_totals_per_year[payer][i] for payer in all_payers)
            ax.annotate(f'{int(total):,}'.replace(',', ' '),
                        xy=(year, total),
                        xytext=(0, 3),
                        textcoords="offset points",
                        ha='center', va='bottom', fontweight='bold')

        ax.set_title("Сравнение затрат по плательщикам по годам", fontsize=16, fontweight='bold')
        ax.set_xlabel("Год", fontsize=12)
        ax.set_ylabel("Сумма (руб.)", fontsize=12)
        ax.legend(title="Плательщики", bbox_to_anchor=(1.05, 1), loc='upper left')
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f'{int(x):,}'.replace(',', ' ')))
        ax.set_xticks(years)
        ax.set_xticklabels([str(y) for y in years])
        ax.grid(axis='y', linestyle='--', alpha=0.7)
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.set_facecolor('#f8f9fa')

        manager = plt.get_current_fig_manager()
        try:
            manager.window.state('zoomed')
        except:
            pass
        plt.tight_layout()
        plt.show()

    def _show_yearly_detail_chart(self, data):
        yearly_payer_totals = {}
        for record in data:
            date_str = record.get("Дата", "").strip()
            sum_str = record.get("Сумма", "").strip()
            payer = record.get("Плательщик", "").strip()
            if not payer:
                continue
            year = None
            if re.match(r'\d{2}\.\d{2}\.\d{4}', date_str):
                try:
                    year = datetime.strptime(date_str, "%d.%m.%Y").year
                except:
                    continue
            try:
                sum_val = float(sum_str.replace(" ", "").replace(",", "."))
            except:
                continue
            if year not in yearly_payer_totals:
                yearly_payer_totals[year] = {}
            if payer not in yearly_payer_totals[year]:
                yearly_payer_totals[year][payer] = 0.0
            yearly_payer_totals[year][payer] += sum_val

        if not yearly_payer_totals:
            messagebox.showwarning("Предупреждение", "Не удалось извлечь данные для графика!")
            return

        years = sorted(yearly_payer_totals.keys())
        n_years = len(years)
        cols = 2
        rows = (n_years + 1) // cols
        fig, axes = plt.subplots(rows, cols, figsize=(16, 6 * rows))
        if n_years == 1:
            axes = [axes]
        elif rows == 1:
            axes = axes if isinstance(axes, np.ndarray) else [axes]
        else:
            axes = axes.flatten()

        all_payers = sorted(set(payer for year_data in yearly_payer_totals.values() for payer in year_data.keys()))
        colors = plt.cm.tab20(np.linspace(0, 1, len(all_payers)))
        payer_colors = {payer: colors[i] for i, payer in enumerate(all_payers)}

        for i, year in enumerate(years):
            ax = axes[i] if n_years > 1 else axes[0]
            year_data = yearly_payer_totals[year]
            sorted_payers = sorted(year_data.items(), key=lambda x: x[1], reverse=True)
            payers = [item[0] for item in sorted_payers]
            amounts = [item[1] for item in sorted_payers]
            year_colors = [payer_colors[payer] for payer in payers]
            bars = ax.bar(payers, amounts, color=year_colors)
            for bar in bars:
                height = bar.get_height()
                ax.annotate(f'{int(height):,}'.replace(',', ' '),
                            xy=(bar.get_x() + bar.get_width() / 2, height),
                            xytext=(0, 3),
                            textcoords="offset points",
                            ha='center', va='bottom', fontweight='bold')
            ax.set_title(f"Затраты по плательщикам за {year} год", fontsize=14, fontweight='bold')
            ax.set_ylabel("Сумма (руб.)", fontsize=12)
            ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f'{int(x):,}'.replace(',', ' ')))
            ax.tick_params(axis='x', rotation=45)
            ax.grid(axis='y', linestyle='--', alpha=0.7)
            ax.spines['top'].set_visible(False)
            ax.spines['right'].set_visible(False)
            ax.set_facecolor('#f8f9fa')

        for j in range(n_years, len(axes)):
            if n_years > 1:
                axes[j].set_visible(False)

        manager = plt.get_current_fig_manager()
        try:
            manager.window.state('zoomed')
        except:
            pass
        plt.tight_layout()
        plt.show()

    def _show_monthly_chart(self, data):
        """График по месяцам за текущий год."""
        current_year = datetime.now().year
        monthly_totals = {i: 0.0 for i in range(1, 13)}

        for record in data:
            date_str = record.get("Дата", "").strip()
            sum_str = record.get("Сумма", "").strip()
            if not re.match(r'\d{2}\.\d{2}\.\d{4}', date_str):
                continue
            try:
                dt = datetime.strptime(date_str, "%d.%m.%Y")
                if dt.year != current_year:
                    continue
                sum_val = float(sum_str.replace(" ", "").replace(",", "."))
                monthly_totals[dt.month] += sum_val
            except:
                continue

        months = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн",
                  "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]
        totals = [monthly_totals[i] for i in range(1, 13)]

        fig, ax = plt.subplots(figsize=(14, 7))
        bars = ax.bar(months, totals, color='teal')
        for bar in bars:
            height = bar.get_height()
            if height > 0:
                ax.annotate(f'{int(height):,}'.replace(',', ' '),
                            xy=(bar.get_x() + bar.get_width() / 2, height),
                            xytext=(0, 3),
                            textcoords="offset points",
                            ha='center', va='bottom', fontweight='bold')

        ax.set_title(f"Затраты по месяцам за {current_year} год", fontsize=16, fontweight='bold')
        ax.set_xlabel("Месяц", fontsize=12)
        ax.set_ylabel("Сумма (руб.)", fontsize=12)
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, _: f'{int(x):,}'.replace(',', ' ')))
        ax.grid(axis='y', linestyle='--', alpha=0.7)
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.set_facecolor('#f8f9fa')

        manager = plt.get_current_fig_manager()
        try:
            manager.window.state('zoomed')
        except:
            pass
        plt.tight_layout()
        plt.show()

    def on_exit(self):
        self.root.destroy()


def main():
    root = tk.Tk()
    app = RegistrumApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()